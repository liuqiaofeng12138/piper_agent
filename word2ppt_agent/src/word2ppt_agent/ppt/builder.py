from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from word2ppt_agent.config.loader import Word2PPTConfig


@dataclass
class BuildResult:
    slide_count: int
    output_path: Path


def _normalize_lines(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if line:
            lines.append(line)
    if not lines:
        chunks = re.split(r"\n{2,}", text.strip())
        lines = [c.strip() for c in chunks if c.strip()]
    return lines


def _chunk_bullets(lines: list[str], cfg: Word2PPTConfig) -> list[list[str]]:
    bullets: list[str] = []
    for line in lines:
        if len(line) <= cfg.max_chars_per_bullet:
            bullets.append(line)
            continue
        start = 0
        while start < len(line):
            bullets.append(line[start : start + cfg.max_chars_per_bullet])
            start += cfg.max_chars_per_bullet

    slides: list[list[str]] = []
    buf: list[str] = []
    for b in bullets:
        buf.append(b)
        if len(buf) >= cfg.max_bullets_per_slide:
            slides.append(buf)
            buf = []
    if buf:
        slides.append(buf)
    return slides


def build_presentation(
    documents: list[tuple[str, str]],
    title: str,
    output_path: Path,
    cfg: Word2PPTConfig,
) -> BuildResult:
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title or "文档内容汇总"
    subtitle = slide.placeholders[1]
    names = "、".join(n for n, _ in documents)
    subtitle.text = f"共 {len(documents)} 个文件：{names}" if names else "上传文件内容汇总"

    slide_count = 1
    content_layout = prs.slide_layouts[1]

    def add_bullet_slide(slide_title: str, group: list[str]) -> None:
        nonlocal slide_count
        s = prs.slides.add_slide(content_layout)
        s.shapes.title.text = slide_title
        body = s.shapes.placeholders[1].text_frame
        body.clear()
        for i, item in enumerate(group):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = item
            p.level = 0
            p.font.size = Pt(18)
        slide_count += 1

    for filename, text in documents:
        if slide_count >= cfg.max_slides_per_doc * len(documents) + 2:
            break
        lines = _normalize_lines(text)
        bullet_groups = _chunk_bullets(lines, cfg)
        if not bullet_groups:
            add_bullet_slide(Path(filename).name, ["（未能提取到可展示的文本要点）"])
            continue
        stem = Path(filename).stem
        for idx, group in enumerate(bullet_groups[: cfg.max_slides_per_doc]):
            title = stem if idx == 0 else f"{stem}（续 {idx}）"
            add_bullet_slide(title, group)

    prs.core_properties.author = "word2ppt-agent"
    prs.core_properties.title = title
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return BuildResult(slide_count=slide_count, output_path=output_path)


def default_ppt_title(user_message: str, doc_names: list[str]) -> str:
    msg = (user_message or "").strip()
    if msg and len(msg) <= 60 and "附件" not in msg:
        return msg
    if len(doc_names) == 1:
        return f"{Path(doc_names[0]).stem} — 内容汇总"
    return "多文档内容汇总"
